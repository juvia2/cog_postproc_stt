"""Extract source observations from the Experiment 1 workbook without scoring them."""

import hashlib
import json
import sys
import warnings
from pathlib import Path

import openpyxl


def main() -> None:
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    warnings.filterwarnings("ignore", message="Sparkline Group extension")
    workbook = openpyxl.load_workbook(source, data_only=False)
    models = workbook["Model 인식 결과"]

    cases = []
    outputs = []
    model_names = ["A.", "Gemini", "CLOVA"]

    for case_number in range(1, 51):
        case_id = f"Case_{case_number:03d}"
        first_model_row = 2 + (case_number - 1) * 3
        source_utterances = [models.cell(first_model_row + offset, 2).value for offset in range(3)]
        source_notes = [models.cell(first_model_row + offset, 3).value for offset in range(3)]
        if len(set(source_utterances)) != 1 or len(set(source_notes)) != 1:
            raise ValueError(f"Case values differ within {case_id}")
        cases.append(
            [
                case_id,
                source_utterances[0],
                source_notes[0],
                first_model_row,
                first_model_row + 2,
                "원문 발화의 실제 음성 파일 미확인",
            ]
        )
        for offset, model_name in enumerate(model_names):
            row_number = first_model_row + offset
            outputs.append(
                [
                    case_id,
                    model_name,
                    models.cell(row_number, 5).value,
                    models.cell(row_number, 6).value,
                    row_number,
                    "과거 인식 기록; 새 음성 재녹음 후 다시 산출 필요",
                ]
            )

    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    payload = {
        "source_path": str(source),
        "source_sha256": digest,
        "cases": cases,
        "model_outputs": outputs,
    }
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "cases": len(cases),
        "model_outputs": len(outputs),
        "model_outputs_missing_text": sum(row[2] is None for row in outputs),
        "source_sha256": digest,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
